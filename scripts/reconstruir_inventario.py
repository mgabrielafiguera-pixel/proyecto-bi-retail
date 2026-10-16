"""Reconstruye el inventario mensual por SKU hasta abril de 2026.

Fuentes (todas anonimizadas):
    02_anonimizado/inventario_sku_mes.csv    fotos reales de cierre de mes, dic-2023 a jun-2025
    02_anonimizado/inventario_marca_mes.csv  resumen por marca, ene-2024 a dic-2025
    03_powerbi/retail_bi_powerbi.xlsx        compras (Sell In) y ventas (Sell Out) por SKU

Método:
    1. Dic-2023 a jun-2025: foto real (Origen = Real). Negativos -> 0.
    2. Jul-2025 a dic-2025: kardex desde la foto de jun-2025
       (IF = IF anterior + compras - ventas, mínimo 0) y calibración del total con
       el índice del resumen por marca (Origen = Estimado calibrado).
    3. Ene-2026 a abr-2026: kardex desde dic-2025, sin ancla (Origen = Estimado kardex).
    Venta_Sin_Stock = Sí cuando el kardex habría quedado negativo: se vendió más de lo
    registrado (recepción no registrada o error de inventario). El stock se deja en 0.

Además corre un backtest: reconstruye ene-jun 2025 desde dic-2024 y lo compara con
las fotos reales.

Uso: python3 scripts/reconstruir_inventario.py
Salidas:
    02_anonimizado/inventario_reconstruido.csv
    03_powerbi/retail_bi_powerbi.xlsx  tabla Inventario (nueva) y SKU sin movimiento en la tabla Producto
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parents[1]
XLSX = RAIZ / "03_powerbi" / "retail_bi_powerbi.xlsx"
P = lambda s: pd.Period(s, "M")
ULTIMA_FOTO, FIN_RESUMEN, FIN = P("2025-06"), P("2025-12"), P("2026-04")


def cargar():
    inv = pd.read_csv(RAIZ / "02_anonimizado" / "inventario_sku_mes.csv", parse_dates=["Fecha"])
    inv["M"] = inv["Fecha"].dt.to_period("M")
    res = pd.read_csv(RAIZ / "02_anonimizado" / "inventario_marca_mes.csv")
    res["M"] = pd.PeriodIndex.from_fields(year=res["Anio"], month=res["Mes"], freq="M")
    prod = pd.read_excel(XLSX, sheet_name="Producto", usecols=["COD_SKU", "CODIGO_INTERNO_MARCA"])
    v = pd.read_excel(XLSX, sheet_name="Ventas", usecols=["COD_SKU", "Cantidad", "Fecha"])
    c = pd.read_excel(XLSX, sheet_name="Compras", usecols=[
        "COD_SKU", "Cantidad", "Fecha_Recepcion", "Es_Posible_Duplicado", "Precio_Cero"])
    # Devoluciones duplicadas con costo erróneo (ver MODELO.md, Calidad de datos)
    c = c[~((c["Es_Posible_Duplicado"] == "Sí") & (c["Precio_Cero"] == "Sí"))]
    si = c.groupby(["COD_SKU", c["Fecha_Recepcion"].dt.to_period("M")])["Cantidad"].sum().unstack(fill_value=0)
    so = v.groupby(["COD_SKU", v["Fecha"].dt.to_period("M")])["Cantidad"].sum().unstack(fill_value=0)
    marca = (pd.concat([prod.rename(columns={"CODIGO_INTERNO_MARCA": "MARCA"}),
                        inv[["COD_SKU", "COD_MARCA"]].rename(columns={"COD_MARCA": "MARCA"})])
               .dropna().groupby("COD_SKU")["MARCA"].first())
    real = inv.groupby(["COD_SKU", "M"])["Unidades"].sum().unstack(fill_value=0).clip(lower=0)
    return inv, res, si, so, marca, real


def kardex(base, meses, si, so):
    """Proyecta el stock por SKU mes a mes. Devuelve (stock, quiebres)."""
    skus = base.index.union(si.index).union(so.index)
    cur = base.reindex(skus, fill_value=0).astype(float)
    stock, quiebre = {}, {}
    for m in meses:
        cur = cur + si.get(m, pd.Series(dtype=float)).reindex(skus, fill_value=0) \
                  - so.get(m, pd.Series(dtype=float)).reindex(skus, fill_value=0)
        quiebre[m] = cur < 0
        cur = cur.clip(lower=0)
        stock[m] = cur.copy()
    return pd.DataFrame(stock), pd.DataFrame(quiebre)


def calibrar(stock, base, meses, ancla, res, marca):
    """Escala el stock de las marcas del resumen para que su total siga el índice del resumen."""
    marcas_res = set(res["COD_MARCA"])
    idx_res = res.groupby("M")["Unidades"].sum()
    en_res = marca.reindex(stock.index).isin(marcas_res)
    total_ancla = base[marca.reindex(base.index).isin(marcas_res)].sum()
    out = stock.copy()
    for m in meses:
        if m in idx_res.index:
            objetivo = total_ancla * idx_res[m] / idx_res[ancla]
            actual = stock.loc[en_res, m].sum()
            if actual > 0:
                out.loc[en_res, m] = stock.loc[en_res, m] * objetivo / actual
    return out.round(0)


def backtest(real, si, so, res, marca):
    ancla, meses = P("2024-12"), pd.period_range("2025-01", "2025-06", freq="M")
    base = real[ancla]
    k, _ = kardex(base, meses, si, so)
    kc = calibrar(k, base, meses, ancla, res, marca)
    obs = real.reindex(k.index, fill_value=0)[meses]
    filas = []
    for nombre, est in [("Kardex", k), ("Kardex calibrado", kc)]:
        err = (est.sum() - obs.sum()) / obs.sum() * 100
        wape = (est - obs).abs().sum() / obs.sum() * 100
        for m in meses:
            filas.append({"Metodo": nombre, "Mes": str(m), "Error_total_%": round(err[m], 1),
                          "WAPE_SKU_%": round(wape[m], 0)})
    return pd.DataFrame(filas)


def reconstruir(real, si, so, res, marca):
    m1 = pd.period_range(ULTIMA_FOTO + 1, FIN_RESUMEN, freq="M")
    k1, q1 = kardex(real[ULTIMA_FOTO], m1, si, so)
    k1 = calibrar(k1, real[ULTIMA_FOTO], m1, ULTIMA_FOTO, res, marca)
    m2 = pd.period_range(FIN_RESUMEN + 1, FIN, freq="M")
    k2, q2 = kardex(k1[FIN_RESUMEN], m2, si, so)

    def largo(df, nombre, origen):
        return df.stack().rename(nombre).reset_index().set_axis(["COD_SKU", "M", nombre], axis=1).assign(Origen=origen)

    partes = [
        largo(real.loc[:, :ULTIMA_FOTO], "Unidades", "Real").assign(Venta_Sin_Stock="No"),
        largo(k1, "Unidades", "Estimado calibrado").merge(
            largo(q1, "Q", "").drop(columns="Origen"), on=["COD_SKU", "M"]),
        largo(k2, "Unidades", "Estimado kardex").merge(
            largo(q2, "Q", "").drop(columns="Origen"), on=["COD_SKU", "M"]),
    ]
    for p in partes[1:]:
        p["Venta_Sin_Stock"] = p.pop("Q").map({True: "Sí", False: "No"})
    out = pd.concat(partes, ignore_index=True)
    # Solo filas con stock o con quiebre; los ceros sin historia no aportan
    out = out[(out["Unidades"] > 0) | (out["Venta_Sin_Stock"] == "Sí")]
    out["Fecha"] = out["M"].dt.to_timestamp(how="end").dt.normalize()
    return out[["COD_SKU", "Fecha", "Unidades", "Origen", "Venta_Sin_Stock"]].sort_values(["Fecha", "COD_SKU"])


def productos_sin_movimiento(inv):
    """Filas nuevas de Producto para los SKUI (stock sin ventas ni compras 2023-2026)."""
    marca = pd.read_excel(XLSX, sheet_name="Marca")
    depto = pd.read_excel(XLSX, sheet_name="Departamento")
    cols = list(pd.read_excel(XLSX, sheet_name="Producto", nrows=0).columns)
    d = inv[inv["COD_SKU"].str.startswith("SKUI")].drop_duplicates("COD_SKU").copy()
    d["Departamento"] = d["Departamento"].replace({"ACCESORIOS S": "ACCESORIOS"})
    cod_dep = depto.set_index("DEPARTAMENTO_ORIGINAL")["COD_DEPARTAMENTO"]
    prov = marca.set_index("CODIGO_INTERNO_MARCA")["CODIGO_INTERNO_PROVEEDOR"]
    filas = pd.DataFrame({
        "COD_SKU": d["COD_SKU"],
        "CODIGO_INTERNO_MARCA": d["COD_MARCA"].fillna("MARC0000"),
        "COD_DEPARTAMENTO": d["Departamento"].map(cod_dep),
        "DEPARTAMENTO": d["Departamento"],
        "en_ventas": "N", "en_compras": "N", "filas_ventas": 0, "filas_compras": 0,
    })
    filas["CODIGO_INTERNO_PROVEEDOR"] = filas["CODIGO_INTERNO_MARCA"].map(prov).fillna("PROV0000")
    return filas.reindex(columns=cols)


def escribir_excel(inventario, nuevos):
    wb = load_workbook(XLSX)
    if "Inventario" in wb.sheetnames:
        del wb["Inventario"]
    ws = wb.create_sheet("Inventario")
    ws.append(list(inventario.columns))
    for fila in inventario.itertuples(index=False):
        ws.append([fila.COD_SKU, fila.Fecha.to_pydatetime(), int(fila.Unidades), fila.Origen, fila.Venta_Sin_Stock])
    for celda in ws["B"][1:]:
        celda.number_format = "yyyy-mm-dd"
    ws.add_table(Table(displayName="Inventario", ref=f"A1:E{ws.max_row}",
                       tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)))
    wp = wb["Producto"]
    existentes = {c.value for c in wp["A"]}
    for fila in nuevos.itertuples(index=False):
        if fila.COD_SKU not in existentes:
            wp.append([None if pd.isna(x) else x for x in fila])
    # La tabla de Excel debe crecer para que Power BI vea las filas nuevas
    wp.tables["Producto"].ref = f"A1:{get_column_letter(wp.max_column)}{wp.max_row}"
    wb.save(XLSX)


def main():
    inv, res, si, so, marca, real = cargar()
    bt = backtest(real, si, so, res, marca)
    print("Backtest ene-jun 2025 (desde la foto de dic-2024):")
    print(bt.pivot(index="Mes", columns="Metodo", values=["Error_total_%", "WAPE_SKU_%"]).to_string())
    inventario = reconstruir(real, si, so, res, marca)
    inventario.to_csv(RAIZ / "02_anonimizado" / "inventario_reconstruido.csv", index=False, date_format="%Y-%m-%d")
    resumen = inventario.groupby(["Fecha", "Origen"]).agg(
        Unidades=("Unidades", "sum"), SKU=("COD_SKU", "nunique"),
        Ventas_Sin_Stock=("Venta_Sin_Stock", lambda s: (s == "Sí").sum()))
    print("\nInventario total por mes:")
    print(resumen.to_string())
    nuevos = productos_sin_movimiento(inv)
    escribir_excel(inventario, nuevos)
    print(f"\nHoja Inventario: {len(inventario):,} filas | SKU nuevos en Producto: {len(nuevos):,}")


if __name__ == "__main__":
    main()
